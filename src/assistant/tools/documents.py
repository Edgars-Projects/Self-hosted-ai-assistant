"""PDF and Excel generation for files the assistant sends back to the user."""

from __future__ import annotations

import json
import os
from typing import Any

from ..config import DL_DIR


def _output_path(filename: str, ext: str) -> str:
    """Resolve a safe path inside the downloads folder with the right extension."""
    os.makedirs(DL_DIR, exist_ok=True)
    if not filename.lower().endswith(ext):
        filename += ext
    return os.path.join(DL_DIR, os.path.basename(filename))


def t_make_pdf(filename: str, title: str, content: str) -> str:
    """Write a simple A4 PDF. Lines starting ``#``/``##``/``###`` become headings
    and ``-``/``*`` lines become bullets."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    path = _output_path(filename, ".pdf")
    styles = getSampleStyleSheet()
    flow: list[Any] = [Paragraph(title, styles["Title"]), Spacer(1, 6 * mm)]
    for block in str(content).split("\n"):
        b = block.strip()
        if not b:
            flow.append(Spacer(1, 3 * mm))
        elif b.startswith("### "):
            flow.append(Paragraph(b[4:], styles["Heading3"]))
        elif b.startswith("## "):
            flow.append(Paragraph(b[3:], styles["Heading2"]))
        elif b.startswith("# "):
            flow.append(Paragraph(b[2:], styles["Heading1"]))
        elif b.startswith(("- ", "* ")):
            flow.append(Paragraph("• " + b[2:], styles["BodyText"]))
        else:
            flow.append(Paragraph(b, styles["BodyText"]))
    SimpleDocTemplate(path, pagesize=A4).build(flow)
    return f"wrote {path} — send it with send_file"


def t_make_xlsx(filename: str, sheets: dict[str, list[list[Any]]] | str) -> str:
    """Write an Excel workbook from ``{"Sheet1": [[row], [row]]}``.

    The first row of each sheet is bolded, columns are auto-sized, and strings
    starting with ``=`` stay as live formulas. ``sheets`` may be a JSON string.
    """
    import openpyxl

    data: dict[str, list[list[Any]]] = json.loads(sheets) if isinstance(sheets, str) else sheets
    path = _output_path(filename, ".xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in data.items():
        ws = wb.create_sheet(str(name)[:31])
        for row in rows:
            ws.append(list(row))
        for c in ws[1]:
            c.font = openpyxl.styles.Font(bold=True)
        for col in ws.columns:
            width = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 50)
    wb.save(path)
    return f"wrote {path} — send it with send_file"
